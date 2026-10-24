"""Build the Assignment 3 report from recorded evidence, never from API calls.

Run analysis/process.py and analysis/figures.py first. The authoring-operation
marker and visual render review are coordinated separately by the operator.
This script never reads .env, invokes a provider, or changes raw observations.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
from pathlib import Path
import re
import statistics
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FILE_NAME = 'RMT_A3_Team_AkhanovYerdaulet_ShadiyarbekSyrym_SabyrzhanAkylbek.docx'
IDENTITY = 'Group SE-2424 | Assignment 3 | Akhanov Yerdaulet, Shadiyarbek Syrym, Sabyrzhan Akylbek'
ACCESS = '24 October 2026'
ATTRIBUTION = '© openrouteservice by HeiGIT | Data from OpenStreetMap'


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def number(value, digits=2):
    if value is None or value == '':
        return 'Not estimable'
    if isinstance(value, str):
        try:
            value = float(value)
        except ValueError:
            return value
    return f'{value:,.{digits}f}'.rstrip('0').rstrip('.') if digits else f'{int(value):,}'


def truth(value):
    return str(value).lower() in {'true', '1', 'yes'}


def set_font(run, size=None, bold=None):
    run.font.name = 'Arial'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_hyperlink(paragraph, url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '174B76'); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = url; run.append(text)
    link.append(run); paragraph._p.append(link)


def text_runs(paragraph, text, size=None):
    for piece in re.split(r'(https?://[^\s]+)', str(text)):
        if piece.startswith(('https://', 'http://')):
            add_hyperlink(paragraph, piece)
        elif piece:
            set_font(paragraph.add_run(piece), size=size)


class Writer:
    def __init__(self):
        self.doc = Document()
        self.prose = []
        self.figures = []
        self.table_count = 0
        section = self.doc.sections[0]
        section.page_width = Cm(21); section.page_height = Cm(29.7)
        section.top_margin = Cm(2); section.bottom_margin = Cm(2)
        section.left_margin = Cm(2); section.right_margin = Cm(2)
        section.header_distance = Cm(.8); section.footer_distance = Cm(.8)
        for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
            style = self.doc.styles[name]
            style.font.name = 'Arial'; style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.line_spacing = 1.08
        # The bundled Word defaults can attach a coloured rule to Title.
        # This report uses plain black heading text without decorative rules.
        for style in self.doc.styles:
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        for name, size in [('Title', 16), ('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 11)]:
            style = self.doc.styles[name]; style.font.size = Pt(size); style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(12)
        self.doc.styles['Caption'].font.size = Pt(9)
        self.doc.styles['Caption'].font.bold = False
        self.doc.styles['Caption'].paragraph_format.keep_with_next = False
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(footer.add_run('SE-2424 | Assignment 3 | '), 8)
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
        update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true')
        self.doc.settings.element.append(update)

    def p(self, text, count=True, style=None):
        p = self.doc.add_paragraph(style=style)
        for border in list(p._p.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        text_runs(p, text)
        p.paragraph_format.widow_control = True
        if count:
            self.prose.append(text)
        return p

    def h(self, text, level=2):
        self.doc.add_heading(text, level=level)

    def table(self, headers, rows, widths=None):
        self.table_count += 1
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if widths is None:
            widths = [17 / len(headers)] * len(headers)
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        props = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'B8B8B8'); borders.append(el)
        props.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for side in ('top', 'bottom', 'left', 'right'):
            el = OxmlElement('w:' + side); el.set(qn('w:w'), '80'); el.set(qn('w:type'), 'dxa'); margins.append(el)
        props.append(margins)
        header = table.rows[0]
        repeat = OxmlElement('w:tblHeader'); header._tr.get_or_add_trPr().append(repeat)
        all_rows = [headers] + list(rows)
        keep_whole = len(all_rows) <= 10
        for i, values in enumerate(all_rows):
            row = header if i == 0 else table.add_row()
            cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
            for j, value in enumerate(values):
                cell = row.cells[j]; cell.width = Cm(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1
                # Keep short evidence tables together. LibreOffice 26.2 can omit
                # continuation rows of a linked two-column table at a page edge.
                p.paragraph_format.keep_with_next = (i < len(all_rows) - 1) if keep_whole else (i == 0)
                text_runs(p, '' if value is None else value, size=9)
                if i == 0:
                    fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'EAEAEA'); cell._tc.get_or_add_tcPr().append(fill)
                    for run in p.runs:
                        run.bold = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return table

    def figure(self, root, entry, index):
        path = root / entry['file']
        contexts = {
            'analysis_request_quality.png': 'distinguishes successful requests from recorded failures',
            'analysis_period_coverage.png': 'checks which predefined periods actually received usable observations',
            'analysis_provider_durations.png': 'compares recorded control estimates across the fixed route identifiers',
            'analysis_distance_duration.png': 'checks the relationship between returned distance and estimated duration',
            'analysis_request_latency.png': 'shows the request-time evidence used for the instrument latency expectation',
            'analysis_recorded_extract.png': 'preserves a readable extract of the actual observation file',
            'route_map.png': 'maps saved ORS paths for the fixed endpoints; these are modeled routes rather than tracked journeys',
        }
        context = entry.get('context') or contexts.get(path.name, 'documents the team-operated instrument and its recorded execution')
        reference = self.p(f'Figure {index} {context}.', count=True)
        reference.paragraph_format.keep_with_next = True
        p = self.doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = p.add_run().add_picture(str(path), width=Cm(16.6))
        shape._inline.docPr.set('descr', entry['caption'])
        shape._inline.docPr.set('title', f'Figure {index}')
        source_label = ', '.join(entry['sources']) if len(entry['sources']) <= 3 else entry['source_manifest']
        credit = entry.get('credit', ATTRIBUTION if entry.get('category') != 'evidence' else 'Team-produced execution evidence')
        self.p(f'Figure {index}. {entry["caption"]} Source: {source_label}. {credit}.', count=False, style='Caption')
        self.figures.append({'number': index, **entry, 'sha256': digest(path)})


def evaluate_expectations(expected, observations, summary, config, manual_cycles, cron_cycles):
    unique = [r for r in observations if r.get('duplicate_kind') != 'exact']
    ors = [r for r in unique if r['provider'] == 'ors']
    pilot = [r for r in ors if r['phase'] == 'pilot']
    pilot_cycles = len({r['cycle_id'] for r in pilot})
    pilot_routes = len({r['route_id'] for r in pilot if truth(r['usable'])})
    n = len(ors)
    invalid = sum(not truth(r['usable']) for r in ors)
    fast = sum(r['latency_s'] != '' and float(r['latency_s']) <= 10 for r in ors)
    scheduled = sum(r['phase'] == 'scheduled' for r in ors)
    manual_verifications = sum(r['phase'] == 'scheduled' and r['cycle_id'] in manual_cycles for r in ors)
    verified_cron = sum(r['phase'] == 'scheduled' and r['cycle_id'] in cron_cycles for r in ors)
    unknown_trigger = scheduled - manual_verifications - verified_cron
    window_ended = datetime.fromisoformat(summary['generated_at_utc']) >= datetime.fromisoformat(config['collection_end_utc'])
    output = []
    for entry in expected['expected_results']:
        eid = entry['id']; verdict = 'cannot judge'
        evidence = 'results/summary.json; results/pairs.csv'
        finding = f'{summary["n_primary_pairs"]} primary pairs; no permitted traffic-aware comparison.'
        if eid == 'ER5':
            finding = f'{invalid}/{n} unique ORS requests failed the valid-route screen ({number(100*invalid/n) if n else "not estimable"}%).'
            if n:
                verdict = 'trending toward' if invalid/n < .02 else 'trending against'
            evidence = 'results/observations.csv; results/request_quality.csv'
        elif eid == 'ER6':
            finding = f'{fast}/{n} unique ORS requests finished within 10 seconds ({number(100*fast/n) if n else "not estimable"}%).'
            if n:
                verdict = 'trending toward' if fast/n >= .95 else 'trending against'
            evidence = 'results/observations.csv; figures/analysis_request_latency.png'
        elif eid == 'ER7':
            finding = f'{pilot_routes}/{config["route_count"]} routes have a usable pilot response; {pilot_cycles}/{config["pilot_cycles"]} pilot cycles are represented.'
            if pilot_routes == config['route_count']:
                verdict = 'trending toward'
            elif pilot_cycles >= config['pilot_cycles']:
                verdict = 'trending against'
            evidence = 'results/observations.csv; figures/analysis_provider_durations.png'
        elif eid == 'ER8':
            finding = f'{verified_cron} verified cron-triggered attempts; {manual_verifications} manual workflow_dispatch verification records (collector phase scheduled); collection window {"ended" if window_ended else "still open"}.'
            if unknown_trigger:
                finding += f' Trigger unverified for {unknown_trigger} further records.'
            if window_ended and not unknown_trigger:
                verdict = 'trending toward' if verified_cron >= 692 else 'trending against'
            evidence = 'results/observations.csv; docs/cloud_run.json; config/study.json'
        output.append({'id': eid, 'expectation': entry['expectation'], 'observed': finding, 'verdict': verdict, 'evidence': evidence})
    return output


def build(root, output):
    required = ['config/study.json', 'docs/expected_results.json', 'docs/references.json',
                'docs/api_check.md', 'docs/research_log.md', 'data/routes.csv',
                'results/summary.json', 'results/observations.csv', 'results/data_dictionary.csv',
                'results/cleaning_log.csv', 'results/coverage.csv', 'results/pairs.csv',
                'figures/analysis_manifest.json', 'figures/route_map_manifest.json', 'collector/collect.py', 'analysis/process.py']
    if (root / 'docs/reproduce_report.md').is_file():
        required.append('docs/reproduce_report.md')
    if (root / 'docs/cloud_run.json').is_file():
        required.append('docs/cloud_run.json')
    missing = [p for p in required if not (root / p).is_file()]
    if missing:
        raise SystemExit('Required evidence absent: ' + ', '.join(missing))
    config = read_json(root / 'config/study.json')
    expected = read_json(root / 'docs/expected_results.json')
    summary = read_json(root / 'results/summary.json')
    refs = read_json(root / 'docs/references.json')
    observations = read_csv(root / 'results/observations.csv')
    dictionary = read_csv(root / 'results/data_dictionary.csv')
    cleaning = read_csv(root / 'results/cleaning_log.csv')
    coverage = read_csv(root / 'results/coverage.csv')
    routes = read_csv(root / 'data/routes.csv')
    manifest = read_json(root / 'figures/analysis_manifest.json')
    figures = [entry for entry in manifest['figures'] if entry['status'] == 'created']
    analysis_order = ['analysis_request_quality.png', 'analysis_period_coverage.png',
                      'analysis_provider_durations.png', 'analysis_distance_duration.png',
                      'analysis_request_latency.png', 'analysis_recorded_extract.png']
    figures.sort(key=lambda e: analysis_order.index(Path(e['file']).name) if Path(e['file']).name in analysis_order else 100)
    for manifest_path, category in [('figures/route_map_manifest.json', 'map'), ('figures/evidence_manifest.json', 'evidence')]:
        if not (root / manifest_path).exists():
            continue
        supplemental = read_json(root / manifest_path)
        entries = supplemental if isinstance(supplemental, list) else supplemental['figures']
        for original in entries:
            if original.get('status', 'created') != 'created':
                continue
            entry = dict(original)
            entry['file'] = entry.get('path') or entry['file']
            entry['sources'] = entry.get('sources') or entry.get('source_files') or ['docs/research_log.md']
            entry['source_manifest'] = manifest_path
            entry['category'] = category
            figures.append(entry)
        if manifest_path not in required:
            required.append(manifest_path)
    if len(figures) < 5:
        raise SystemExit('At least five real study figures are required; run the analysis figure generator.')
    if len(refs) != 10 or any(not r.get('citation') or not r.get('free_url') for r in refs):
        raise SystemExit('The ten Assignment 2 references and public full-text links are required.')
    if summary['n_raw_rows'] != len(observations):
        raise SystemExit('Summary and observations disagree; rerun analysis on the same snapshot.')
    if config['active_providers'] != ['ors'] or summary['n_primary_pairs']:
        raise SystemExit('This report narrative is for the ORS-only instrument pilot. Review it before enabling a primary comparison.')
    if len(config['rq'].split()) > 25:
        raise SystemExit('Research question exceeds 25 words.')
    for entry in figures:
        if not (root / entry['file']).is_file():
            raise SystemExit('Missing figure: ' + entry['file'])
    manual_cycles, cron_cycles = set(), set()
    if (root / 'docs/cloud_run.json').is_file():
        cloud_run = read_json(root / 'docs/cloud_run.json')
        run_match = re.search(r'/actions/runs/(\d+)(?:/|$)', cloud_run.get('url', ''))
        if run_match:
            cycle_id = 'gh-' + run_match.group(1)
            if cloud_run.get('event') == 'workflow_dispatch':
                manual_cycles.add(cycle_id)
            elif cloud_run.get('event') == 'schedule':
                cron_cycles.add(cycle_id)
    er = evaluate_expectations(expected, observations, summary, config, manual_cycles, cron_cycles)
    repo = config['repository_url'].rstrip('/')
    public = lambda path: repo + '/blob/main/' + path
    tree = lambda path: repo + '/tree/main/' + path
    ors_rows = [r for r in observations if r['provider'] == 'ors' and r['duplicate_kind'] != 'exact']
    pilot = [r for r in ors_rows if r['phase'] == 'pilot']
    usable = [r for r in ors_rows if truth(r['usable'])]
    pilot_n = len(pilot)
    scheduled_n = sum(r['phase'] == 'scheduled' for r in ors_rows)
    manual_cloud_n = sum(r['phase'] == 'scheduled' and r['cycle_id'] in manual_cycles for r in ors_rows)
    cron_n = sum(r['phase'] == 'scheduled' and r['cycle_id'] in cron_cycles for r in ors_rows)
    valid_pilot_routes = len({r['route_id'] for r in pilot if truth(r['usable'])})
    periods = Counter(r['period'] for r in usable)
    local_dates = sorted({r['date_astana'] for r in usable})
    first = summary['first_request_utc'] or 'No recorded request'
    last = summary['last_request_utc'] or 'No recorded request'
    prereg_sha = digest(root / 'docs/expected_results.json')
    w = Writer()
    w.p(IDENTITY, count=False)
    w.p('Astana route estimate methodology', count=False, style='Title')
    w.p('Status: executed ORS instrument pilot; the primary traffic-aware comparison has not been performed.', count=False)
    w.p(f'Evidence snapshot: {summary["generated_at_utc"]}. Public repository: {repo}', count=False)

    w.h('Segment 1. Methodology', 1)
    w.h('T1 Declare basic information')
    w.p(config['rq'], count=False)
    w.p('The question isolates disagreement between services for common endpoints and time periods. Its strength is an observable comparison with a fixed unit and a clear temporal contrast. Its weakness is that services can recommend different streets and apply undisclosed models. Neither estimate is an observed journey. The original Assignment 2 Google Maps–2GIS design was amended after access and publication checks. TomTom–2GIS remains the prospective question, while the completed work concerns the ORS instrument. This distinction prevents successful software operation from being presented as an answer about congestion.')
    drive_status = 'Archive upload verified' if config.get('drive_archive_uploaded') else 'Archive upload pending verification at this snapshot'
    w.table(['Item', 'Entry'], [['Group and team', config['group'] + ' | ' + config['team']], ['Repository', repo], ['Drive folder', config['drive_url'] + '\n' + drive_status], ['Folders', 'collector/: extraction; config/: fixed design; data/routes.csv: sample; data/raw/: original observations; analysis/: processing; results/: working data; figures/: evidence; report/: reproducible document; docs/: provenance and journal']], [3.4, 13.6])
    w.h('T2 Record your Approach Card')
    w.p('The approach is quantitative secondary-data analysis using repeated observations. Researchers select requests but do not manipulate road conditions, allocate traffic treatments or recruit drivers. It is therefore an observational comparison, not an experiment. Repeated endpoints reduce spatial inconsistency and make variation auditable. However, purposive route selection and a short observation window limit generalisation. Repeated requests on one route are dependent; their number cannot be treated as the number of independent roads. ORS offers a practical control for the collection workflow but lacks the live-traffic signal required by the main question.')
    w.table(['Dimension', 'Declared approach'], [['Data and purpose', 'Numerical API estimates; descriptive agreement planned; instrument feasibility observed'], ['Design', 'Observational repeated measures; secondary data; prospective paired comparison'], ['Strength', 'Fixed OD identifiers, timestamps and preserved failures support reproduction'], ['Weakness', 'No ground truth; proprietary provider behavior; dependent routes and short duration']], [3.4, 13.6])
    w.h('T3 Concept table')
    w.table(['Concept', 'Operational definition', 'Measure and unit'], [['Travel-time disagreement', 'For paired estimates A and B, d=A−B and m=(A+B)/2', 'd and |d| in seconds; 100d/m and 100|d|/m in percent'], ['Traffic period', 'Classify request timestamp in Asia/Almaty; start included, end excluded', 'Nominal: peak, offpeak, other'], ['Route comparability', '100|LA−LB|/((LA+LB)/2) ≤ 5; length proxy only', 'Ratio: percent; nominal: comparable yes/no'], ['Instrument reliability', 'Valid route responses and elapsed request time', 'Ratio: success proportion, latency seconds']], [3.1, 8.8, 5.1])
    w.h('T4 What exactly was observed')
    w.table(['Observation element', 'Entry'], [['Unit', 'One provider response for a fixed OD pair at a recorded request start'], ['Observed', 'ORS estimate, distance, geometry, status, timestamps and elapsed request time'], ['Not observed', 'Actual trips, driver behavior, live traffic or TomTom–2GIS disagreement'], ['Comparison unit planned', 'Same route and cycle, provider request starts ≤ 60 seconds apart; same period and date']], [4.3, 12.7])
    w.h('T5 Define your Research strategy')
    w.p('The strategy is a bounded Astana case study with purposive sampling of public arterial corridors and landmarks. Common endpoints specify the requested journey rather than forcing a common returned path. A practical consequence is that analysis must preserve both all eligible pairs and a length-comparable subset. Similar distance alone cannot prove street-level equivalence. The design therefore keeps geometry where its licence permits, records route definitions separately, and avoids interpreting differences as provider error. The instrument phase can proceed independently without changing the predeclared traffic hypotheses.')
    w.h('T6 Define sample, dataset or cases')
    w.table(['Item', 'Entry'], [['Population or case set', 'Fixed public OD pairs on selected Astana arterial corridors; not all city trips'], ['Planned sample', f'{len(routes)} routes; {config["pilot_cycles"]} full pilot cycles; {config["planned_scheduled_requests_per_provider"]} nominal scheduled ORS requests'], ['Achieved snapshot', f'{summary["n_raw_rows"]} raw rows; {summary["n_unique_requests"]} unique requests; {summary["n_usable_observations"]} usable observations; {summary["n_primary_pairs"]} primary pairs'], ['Selection strategy', 'Purposive route and endpoint selection against public-road and corridor criteria'], ['Source', public('data/routes.csv') + '\n' + tree('data/sources')]], [4.1, 12.9])
    w.p('Coordinates were fixed before research collection using public OpenStreetMap endpoint records. Reverse directions are distinct cases because access and road direction can differ. The expected-length column is a spherical straight-line lower bound, not a measured driving distance. It supports a geographic plausibility check and must not be substituted for the returned route length. The sample contains named corridors rather than personal origins. Selection does not give every Astana journey a known probability of inclusion, so the intended conclusions remain conditional on these cases.')
    w.h('T7 Procedure actually followed')
    w.p(f'First, the team reviewed the earlier assignments, checked credential access with public example coordinates and separated those tests from research data. Second, it inspected provider documentation and publication terms. This led to disabling the traffic-provider adapters before the research collection. Third, route definitions, period rules and expectations were fixed on {expected["fixed_date"]}. Fourth, the ORS collector recorded the pilot requests using a schema that also records unsuccessful attempts. Fifth, the processing script produced an audited working file, coverage tables and the original figures below. The journal identifies the operator and any later changes.')
    w.p(f'The pilot plan uses {config["pilot_cycles"]} complete cycles with {config["pilot_pause_seconds"]} seconds between the end of one cycle and the next. The scheduled plan rotates {config["routes_per_scheduled_run"]} routes every {config["scheduled_interval_minutes"]} minutes within a bounded {config["collection_hours"]}-hour window. Each route therefore has a nominal {config["scheduled_interval_minutes"]*len(routes)//config["routes_per_scheduled_run"]}-minute revisit. Scheduler delays remain visible through actual request times. No past observation is backfilled. Reproduction uses the versioned configuration and dependency file; Python {summary["versions"]["python"]} processed this snapshot. A future provider pair must use the same endpoints, car mode and closely aligned request starts.')
    for index, entry in enumerate(figures, 1):
        w.figure(root, entry, index)
    w.h('T8 Ethics, privacy and data handling')
    w.p('No human participant supplies a personal journey or identity, so participant consent is not applicable to this public-coordinate instrument. API access still has contractual conditions. Technical success was insufficient to authorize a public comparison: TomTom requires written approval for benchmarking, while the applicable 2GIS Kazakhstan demo redistribution terms remain unresolved. ORS results are shared under current CC BY-SA terms with attribution. Secrets remain outside version control. The team retains failed requests and exclusion reasons to make limitations visible, while avoiding account identifiers and credential-bearing URLs in public evidence.')
    w.table(['Item', 'Entry'], [['Consent', 'Not applicable: no recruited people or personal journeys'], ['Anonymization', 'Public route identifiers only; no individual tracks'], ['Not collected', 'Names of drivers, home addresses, device identifiers or actual travel histories'], ['Storage, access and retention', 'Private local secrets; public licensed research files and code. Public links must remain available through exam weeks.'], ['Accepted risk', 'Model estimates may be incomplete or inconsistent; records describe agreement and instrument behavior, not navigation advice']], [4.4, 12.6])
    w.h('T9 Expected Results')
    w.p('The expectation record is fixed before the first research observation. Earlier credential checks established access only and did not supply Astana outcomes. Expectations about provider disagreement remain unchanged when collection becomes restricted. Instrument expectations are separately falsifiable through valid-response counts, latency and route coverage. The six scenarios describe possible outcomes rather than simulated findings. A later amendment must carry its own date and reason; it cannot silently replace the existing thresholds after results become visible. Objectives retain the earlier project sequence, which Assignment 2 did not separately renumber.')
    w.table(['No.', 'Expected result', 'Objective', 'Falsification', 'Confidence and reason'], [[e['id'], e['expectation'], e['objective'], e['falsification'], e['confidence']] for e in expected['expected_results']], [1, 5.4, 2.5, 3.8, 4.3])
    w.table(['Outcome', 'Possible scenario'], [[s['level'], s['description']] for s in expected['scenarios']], [2.5, 14.5])
    w.p(f'Fixed date: {expected["fixed_date"]}. Timestamp: {expected["fixed_at_utc"]}. Expected-results file SHA-256: {prereg_sha}. Preregistration commit: b4614355903efdc800f870d2678d2416e0490f81.', count=False)
    w.h('T10 State your method-related expectations')
    w.table(['Method', 'Declaration before analysis'], [['Data analytics / secondary data', 'Primary expectation: larger median absolute disagreement during evening peak. Instrument expectations: <2% invalid routes, ≥95% requests within 10 s, all 20 pilot routes represented. Missing optional baselines, failed responses, duplicates and implausible implied speeds are recorded. API agreement cannot establish actual travel-time accuracy.']], [4.3, 12.7])

    w.h('Segment 2. Data Collection Instrument', 1)
    w.h('T1 Provenance information')
    w.p('The design adapts repeated route and time sampling from Elnaggar et al. (2025) and fixed OD comparison from Fu et al. (2023). It does not reproduce their cities, datasets or validation claims. Hodson (2022) motivates giving each difference measure a defined role. The collector is self-developed because synchronized requests, explicit failures, route identifiers and permission controls must share one audit trail. Its command-line interface writes CSV and licensed response files; it has no participant frontend or personal-data database. Confidence comes from recorded pilot outputs and deterministic validation, not from calling the software a methodology.')
    w.table(['Element', 'Source', 'Status', 'Taken', 'Changed or justified'], [['Overall design', 'Elnaggar et al. (2025); Fu et al. (2023)', 'Adapted', 'Repeated route/time observations; common OD comparisons', 'Astana cases; explicit disagreement target; ORS-only pilot'], ['Sampling', 'Elnaggar et al. (2025)', 'Adapted', 'Defined urban corridors and departure periods', 'Public endpoints; rotating bounded schedule'], ['Instrument', 'Team; informed by Fu et al. (2023)', 'Self-developed', 'Programmatic route estimates', 'Schema audit, failures, quotas and publication gate'], ['Scales or items', 'Not applicable', 'Not applicable', 'No questionnaire or psychometric scale', 'Numerical measures operationalized in Segment 1'], ['Metrics', 'Hodson (2022)', 'Adapted', 'Absolute and squared difference summaries', 'Agreement interpretation; neither provider is truth'], ['Dataset', 'Team collection', 'Self-developed', 'No published dataset copied', 'Public-coordinate ORS estimates with recorded provenance']], [2.1, 3.5, 2.1, 4.3, 5])
    w.h('T2 Replication status')
    w.p('This is an original study informed by earlier methods, rather than a full or partial replication of one published protocol. The team combines its question, route set, collection controls and evidence pipeline. No paper supplies an exact template whose result is being reproduced. Derrow-Pinion et al. (2021) and Soon et al. (2023) provide forecasting context, not evidence that the reviewed providers deploy identical algorithms. The ten-paper Assignment 2 corpus is retained so that design provenance remains traceable without claiming implementation details that those sources do not establish.')
    w.table(['Declaration', 'What was developed'], [['Original study', 'Astana question, purposive sample, auditable collection instrument and descriptive analysis. No single published study was used as the replication template.']], [3.4, 13.6])
    w.h('T3 Your research is original study')
    w.table(['Element', 'Entry'], [['Question or hypothesis', config['rq'] + '\nH1: ' + config['hypothesis_h1'] + '\nH0: ' + config['hypothesis_h0']], ['Methodology and design', 'Observational repeated secondary-data comparison; ORS-only instrument pilot executed'], ['Population, sample and size', f'{len(routes)} fixed Astana OD routes; {len(ors_rows)} unique ORS requests at snapshot'], ['Instrument', 'Versioned collector, route CSV, configuration and raw request records'], ['Technical procedure', 'HTTP requests with timeout and no redirects; recorded errors; immutable raw inputs'], ['Analysis', 'Validation, missingness and descriptive instrument summaries; paired metrics defined prospectively'], ['Main reported result', f'{len(usable)} usable ORS observations; primary traffic comparison not performed']], [4.3, 12.7])
    w.h('T4 Your research is partial replication')
    w.p('Not applicable. No partial-replication claim is made; adaptations are declared in the provenance table.', count=False)

    w.h('Segment 3. Research Evidence', 1)
    w.h('T1 Instrument register')
    w.table(['Instrument', 'Platform', 'Collects or defines', 'Public link'], [['Collector', 'Python CLI', 'Estimates, geometry, status and timing', public('collector/collect.py')], ['Route register', 'CSV', 'Fixed OD coordinates and source links', public('data/routes.csv')], ['Scheduled workflow', 'GitHub Actions', 'Bounded recurring collection and logs', tree('.github/workflows')], ['Processing', 'Python', 'Audited observations and descriptive summaries', public('analysis/process.py')]], [2.3, 2.3, 4.1, 8.3])
    w.h('T2 Log journal and data storage')
    w.p('The storage layout separates original records from the working copy used for interpretation. Raw CSV rows retain source files and request identifiers; processing adds validation decisions instead of overwriting observations. A version manifest records input hashes and software versions. This makes a later report reproducible even when new scheduled records arrive. The research journal records what was done, by whom, what failed and what changed. The current report is a dated snapshot, while the processed observations and documented procedures form the evidence base for Assignment 4.')
    if (root / 'docs/reproduce_report.md').is_file():
        w.p('Exact snapshot reproduction uses frozen inputs rather than later arriving raw files. Follow the saved instructions and input hashes at ' + public('docs/reproduce_report.md'), count=False)
    w.table(['Item', 'Location and purpose'], [['Research journal', public('docs/research_log.md')], ['Raw data', tree('data/raw') + '\nOriginal CSV plus licensed ORS response and geometry files'], ['Working copy', public('results/observations.csv') + '\nAudited observations; raw values retained'], ['Version record', public('results/version_manifest.json')], ['Public repository', repo], ['Cloud folder', config['drive_url'] + '\n' + drive_status]], [3.3, 13.7])
    w.h('T3 Artefacts you developed')
    w.table(['Artefact', 'Study role', 'Public link', 'Evidence'], [['Collector and pilot runner', 'Execute fixed requests and preserve failures', tree('collector'), 'Raw files and research journal'], ['Analysis and figures', 'Audit data and generate original evidence', tree('analysis'), 'Results CSV and numbered figures'], ['Report builder', 'Bind narrative to a reproducible snapshot', public('report/build_report.py'), 'Report snapshot and word count'], ['Protocol and expectations', 'Record rules fixed before research collection', tree('docs'), 'Timestamp, file hash and preregistration commit']], [2.7, 4.3, 6.5, 3.5])
    w.h('T4 Find your block and complete')
    w.p('The selected method is Data analytics / secondary data, so Block E applies. Technical execution supports extraction; it does not turn this study into model training or an experimental treatment comparison.', count=False)
    w.h('E.0 Requirements', 3)
    w.p('For a conceptual road graph G=(V,E) at a fixed traffic snapshot, path cost is T(P,t)=Σ w_e(t), summed over edges e in P; each weight is in seconds. Route selection seeks a permitted low-cost path. This abstraction explains why endpoints alone do not determine an estimate: road access, path choice and time weights may differ. It does not identify TomTom or 2GIS source code. ORS documentation describes a speed model based on mapped road properties; its public service does not supply live traffic. The study evaluates returned quantities without reverse-engineering a proprietary solver.')
    w.p('For an eligible pair, signed difference is A minus B and absolute difference removes its direction. MAE averages absolute differences, whereas RMSE gives larger differences greater influence. Relative difference uses the pair mean as denominator, avoiding a claim that either provider is the reference truth. Quartiles describe spread. Planned uncertainty procedures preserve route clusters rather than treating repeated requests as independent. The pilot receives descriptive summaries only. With no primary pairs, disagreement statistics and peak contrasts remain unestimated, rather than being filled with zeros or inferred from ORS variability.')
    w.table(['Quantity', 'Definition', 'Interpretation'], [['Signed difference', 'd=A−B; seconds', 'Positive means A gives the longer estimate'], ['Absolute difference and MAE', '|d|; MAE=Σ|d|/n', 'Magnitude and average magnitude of disagreement'], ['RMSE', 'sqrt(Σd²/n); seconds', 'Quadratic summary sensitive to large differences'], ['Relative difference / sMAPE', '100d/m; 100|d|/m; m=(A+B)/2', 'Symmetric percentage disagreement'], ['IQR', 'Q3−Q1; linear-interpolated sample quartiles', 'Spread of the stated distribution'], ['Length comparison', '100|LA−LB|/((LA+LB)/2)', '≤5% length agreement is not geometric identity']], [4.2, 6.1, 6.7])
    collector_text = (root / 'collector/collect.py').read_text(encoding='utf-8')
    snippet_start = collector_text.index("    if provider=='ors':", collector_text.index('def extract('))
    snippet_end = collector_text.index("    if provider=='tomtom':", snippet_start)
    snippet = collector_text[snippet_start:snippet_end].rstrip()
    w.p('Extraction excerpt from the saved collector; the complete request construction and error handling are at ' + public('collector/collect.py'), count=False)
    p = w.doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
    r = p.add_run(snippet); r.font.name = 'Consolas'; r.font.size = Pt(9)
    w.h('E.1 Source register', 3)
    w.table(['Source', 'Owner', 'Access date', 'Version', 'Licence / scope', 'Full URL'], [['ORS results', 'HeiGIT', ACCESS, 'Directions v2; build metadata in saved response', 'CC BY-SA 4.0; attribution; active instrument', 'https://account.heigit.org/info/tos'], ['Endpoint facts', 'OpenStreetMap contributors', ACCESS, 'Saved Overpass extract and manifest', 'ODbL; public geographic data', 'https://www.openstreetmap.org/copyright'], ['TomTom docs', 'TomTom', ACCESS, 'Routing v1', 'Documentation only; comparison/storage blocked', 'https://docs.tomtom.com/legal/terms-and-conditions'], ['2GIS docs', '2GIS', ACCESS, 'Routing 7.0.0 global', 'Documentation only; applicable demo rights unresolved', 'https://docs.2gis.com/en/api/navigation/routing/overview'], ['Study records', 'Project team', ACCESS, summary['generated_at_utc'], 'Derived ORS outputs retain CC BY-SA attribution', tree('data')]], [2, 2, 1.8, 3, 4, 4.2])
    w.h('E.2 Data dictionary', 3)
    w.p('The dictionary distinguishes absent optional fields from missing measurements. An empty no-traffic duration is expected for ORS and is not replaced by a fabricated baseline. Empty error and exclusion fields indicate a successful unflagged record, not poor data quality. Missing counts use each named dataset as their denominator. A paired dataset with no records has no observed missingness percentage. The inventory includes every raw, audited and paired column generated by the processing script.', count=True)
    w.table(['Dataset / column', 'Type', 'Unit', 'Considered values or range', 'Missing count / %'], [[r['dataset'] + '\n' + r['column'], r['type'], r['unit'], r['considered_values_or_range'], f'{r["missing_count"]} / ' + (f'{r["missing_pct"]}%' if r['missing_pct'] != '' else 'N/A') + f' (n={r["n_rows"]})'] for r in dictionary], [4.4, 1.7, 2.3, 5.7, 2.9])
    w.h('E.3 Cleaning log', 3)
    w.table(['Step', 'Rule applied', 'Rows before / after', 'What it revealed'], [[r['step'], r['rule'], r['rows_before'] + ' / ' + r['rows_after'], r['finding'] + '\nAffected rows: ' + r['affected_rows']] for r in cleaning], [2.8, 6.9, 2.4, 4.9])
    w.h('E.4 Evidence to present', 3)
    w.table(['Evidence', 'Public location', 'Visible content'], [['Initial extract', tree('data/raw'), 'Original chronological request rows and licensed ORS responses'], ['Processed file', public('results/observations.csv'), 'Source links, values and validation decisions'], ['Codings', public('analysis/process.py'), 'Cleaning, grouping and metric definitions'], ['Figures', public('figures/analysis_manifest.json'), 'Input provenance and captions for original visual evidence']], [3.1, 8.4, 5.5])
    w.h('E.5 Describe the time or category coverage of the final dataset', 3)
    w.p(f'The recorded request interval is {first} to {last} in UTC. Usable observations cover {len(local_dates)} local date(s): {", ".join(local_dates) or "none"}. Classification produces {periods["peak"]} peak, {periods["offpeak"]} off-peak and {periods["other"]} other-period observations. These are coverage counts for the instrument, not evidence of congestion. The period table retains phase and provider so pilot activity cannot masquerade as scheduled coverage. Absent primary pairs prevent any comparison between peak and off-peak disagreement, regardless of how many ORS timestamps exist.')
    w.table(['Phase', 'Local date', 'Period', 'Provider', 'Usable observations'], [[r['phase'] + (' (manual workflow_dispatch verification)' if r['phase'] == 'scheduled' and manual_cloud_n == scheduled_n else ''), r['date_astana'], r['period'], r['provider'], r['n_usable']] for r in coverage], [3, 3.6, 3.5, 3.5, 3.4])

    w.h('Segment 4. Expected Against Observed. Preliminary comparison', 1)
    w.h('T1 Completion table')
    primary_period_text = '; '.join(f'{label} {summary["primary_pairs_by_period"][key]}' for key, label in [('peak', 'Peak'), ('offpeak', 'off-peak'), ('other', 'other')])
    scheduled_display = f'{cron_n} verified cron; {manual_cloud_n} manual workflow_dispatch test records (phase: scheduled)'
    if scheduled_n > cron_n + manual_cloud_n:
        scheduled_display += f'; {scheduled_n-cron_n-manual_cloud_n} trigger unverified'
    w.table(['Data to be collected', 'Planned', 'Collected at snapshot'], [['Route definitions', str(config['route_count']), str(len(routes))], ['ORS pilot requests', str(config['planned_pilot_requests_per_provider']), str(pilot_n)], ['ORS pilot routes with usable response', str(config['route_count']), str(valid_pilot_routes)], ['Scheduled ORS requests', str(config['planned_scheduled_requests_per_provider']) + ' nominal over 48 hours', scheduled_display], ['Primary traffic-provider pairs', 'Prospective; collection disabled pending permissions', str(summary['n_primary_pairs'])], ['Primary peak and off-peak coverage', 'Both required for H1', primary_period_text]], [5.8, 5.6, 5.6])
    w.h('T2 Comparison')
    w.table(['Expected result', 'What actual data shows', 'Verdict', 'Evidence location'], [[e['id'] + ': ' + e['expectation'], e['observed'], e['verdict'], e['evidence']] for e in er], [5.4, 5.2, 2.1, 4.3])
    w.h('T3 Answer in continuous text')
    w.p('The first records test whether the instrument can produce traceable, usable data. They cannot answer the traffic-provider question. The no-traffic field is unavailable as a separate ORS measurement, and all primary disagreement measures lack an eligible provider pair. These are substantive limits of the available instrument, not evidence that traffic effects or provider differences are absent. Pilot success is assessed against the predeclared reliability expectations; it is not a conclusion that one route service is more accurate or that an observed estimate matches a real trip.')
    w.p('Access and terms behaved differently from the initial assumption that a working key would be sufficient for research publication. The amendment on 24 October 2026 consequently restricted active collection to ORS and reduced the scheduled workload to rotating batches. The primary hypotheses, fixed endpoints, time classification and falsification thresholds remain unchanged. Any later traffic-provider collection requires documented permission and a dated activation decision. Earlier ORS observations will retain their control status. This preserves comparability of the instrument record while acknowledging that a permitted, paired traffic sample is still required before the main study can be interpreted.')
    w.p('The earlier literature supports this restraint. Bastarianto et al. (2025) and Valero et al. (2025) distinguish calculated from reported journeys; neither kind can simply be treated as error-free truth. Boeing and Zhou (2026) and Verma et al. (2025) demonstrate useful reference comparisons without removing that distinction. Mansurova et al. (2025) provide local bus data, but its mode and stopping behavior do not validate private-car estimates. Future interpretation must therefore separate model agreement, sampling coverage and independent journey validation instead of using one as a substitute for another.')

    w.h('References', 1)
    bibliography = []
    for ref in refs:
        citation = ref['citation']
        if ref['doi'] not in citation:
            citation += ' https://doi.org/' + ref['doi']
        bibliography.append(citation)
    official = [
        ('2GIS', 'Routing API overview', 'https://docs.2gis.com/en/api/navigation/routing/overview'),
        ('2GIS', 'WebAPI licence offer', 'https://law.2gis.ru/offer-license-agreement-webapi'),
        ('HeiGIT', 'openrouteservice API reference', 'https://giscience.github.io/openrouteservice/api-reference/'),
        ('HeiGIT', 'openrouteservice frequently asked questions', 'https://giscience.github.io/openrouteservice/frequently-asked-questions'),
        ('HeiGIT', 'Terms of service', 'https://account.heigit.org/info/tos'),
        ('HeiGIT', 'Isochrone traffic data: maintainer response', 'https://ask.openrouteservice.org/t/isochrone-trafic-data/7875'),
        ('OpenStreetMap contributors', 'Copyright and licence', 'https://www.openstreetmap.org/copyright'),
        ('TomTom', 'Calculate Route documentation', 'https://docs.tomtom.com/routing-api/documentation/tomtom-maps/v1/calculate-route'),
        ('TomTom', 'Portal terms and conditions', 'https://docs.tomtom.com/legal/terms-and-conditions'),
        ('TomTom', 'Pricing', 'https://docs.tomtom.com/pricing'),
    ]
    for owner, title, url in official:
        bibliography.append(f'{owner} (2026) {title}. Available at: {url} (Accessed: {ACCESS}).')
    for citation in sorted(bibliography, key=str.casefold):
        w.p(citation, count=False)
    w.h('Free-access links', 1)
    for ref in sorted(refs, key=lambda r: r['citation'].casefold()):
        w.p(ref['citation'].split('(')[0].strip() + ' (' + str(ref['year']) + '): ' + ref['free_url'], count=False)
    w.p(f'Project evidence: {repo}\nCloud folder: {config["drive_url"]}', count=False)

    prose_words = len(re.findall(r"\b[\w]+(?:['’−-][\w]+)*\b", '\n'.join(w.prose)))
    if prose_words < 1200:
        raise SystemExit(f'Prose word count {prose_words} is below the required minimum.')
    output.parent.mkdir(parents=True, exist_ok=True)
    w.doc.core_properties.title = 'Assignment 3: Astana route-estimate instrument and methodology'
    w.doc.core_properties.author = config['authors']
    w.doc.core_properties.subject = 'ORS instrument pilot; prospective traffic-provider comparison'
    w.doc.save(output)
    report_dir = root / 'report'; report_dir.mkdir(exist_ok=True)
    evidence_hashes = {path: digest(root / path) for path in required}
    snapshot = {'generated_at_utc': datetime.now(timezone.utc).isoformat(),
                'analysis_snapshot_utc': summary['generated_at_utc'],
                'report_file': output.name, 'report_sha256': digest(output),
                'scope': 'ORS-only instrument pilot; primary comparison not performed',
                'prose_word_count': prose_words,
                'counting_rule': 'Only authored continuous-prose paragraphs; excludes identity, headings, RQ, tables, captions, code, references and links.',
                'figure_count': len(w.figures), 'table_count': w.table_count,
                'versions': {'python': sys.version.split()[0], 'python-docx': importlib.metadata.version('python-docx')},
                'input_sha256': evidence_hashes, 'figures': w.figures,
                'expected_results_comparison': er, 'summary': summary,
                'validation': {'visual_render_review': 'pending separate render review', 'references': len(refs)}}
    (report_dir / 'report_snapshot.json').write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding='utf-8')
    (report_dir / 'report_prose.txt').write_text('\n\n'.join(w.prose), encoding='utf-8')
    (report_dir / 'report_word_count.json').write_text(json.dumps({'prose_words': prose_words, 'required_minimum': 1200, 'exclusions': snapshot['counting_rule']}, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(output), 'prose_words': prose_words, 'figures': len(w.figures), 'tables': w.table_count}))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    build(root, args.output.resolve() if args.output else root / FILE_NAME)


if __name__ == '__main__':
    main()
